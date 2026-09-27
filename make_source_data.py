#!/usr/bin/env python3
"""Build the Source Data workbooks for the MorphoResidual manuscript (Nature Communications).

Run from anywhere:   python make_source_data.py
Writes to ./source_data/ :
    Source_Data_Fig1.xlsx ... Source_Data_Fig5.xlsx   one workbook per main figure, one sheet per panel
    Source_Data_Tables.xlsx                           main Tables 1 and 2
    Supplementary_Tables_long.xlsx                    the three long supplementary tables
    README.md, crosscheck_report.txt

Nothing is imported from the figure scripts.  Every panel table is recomputed with pandas from the SAME
input files the corresponding figure script reads (figures/figdata/*.csv, server_export/pinned/*,
server_export/results/*, review/recalc/*), replicating the arithmetic the script performs on them
(median, held-out median, t-SNE, KM curve, Fisher-z interval, ...).  Where a figure script derives a
plotted quantity, the derivation is repeated here on purpose; if it ever disagrees with the manuscript
text, the assertion/cross-check at the bottom reports it instead of silently drawing a different number.

Sheet layout (all panel sheets):
    row 1  panel title
    row 2  source file(s), relative to the project root
    row 3  column names (units in brackets)
    row 4  one-line explanation of each column
    row 5+ data
Read one back with:  pandas.read_excel(path, sheet_name="Fig1c", header=2, skiprows=[3])
"""
import os
import json
import math
import warnings
import sys
import numpy as np
import pandas as pd
from scipy import stats as st
from scipy.stats import chi2
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from sklearn.manifold import TSNE

warnings.filterwarnings("ignore")

HERE = os.path.dirname(os.path.abspath(__file__))
# the one shared definition of the RPPA antibody rule (a helper module, not a figure script)
sys.path.insert(0, os.path.join(HERE, "figures"))
from rppa_filter import load_rppa  # noqa: E402
FIG = os.path.join(HERE, "figures")
DD = os.path.join(FIG, "figdata")
PIN = os.path.join(HERE, "server_export", "pinned")
RES = os.path.join(HERE, "server_export", "results")
REC = os.path.join(HERE, "review", "recalc")
OUT = os.path.join(HERE, "source_data")
os.makedirs(OUT, exist_ok=True)

COH = ["CCRCC", "LUAD", "UCEC", "GBM", "PDAC"]
ORGAN = {"CCRCC": "Kidney", "LUAD": "Lung", "UCEC": "Uterus", "GBM": "Brain", "PDAC": "Pancreas"}

REL = lambda p: os.path.relpath(p, HERE).replace(os.sep, "/")
CHECKS = []          # (label, computed, expected, tol, ok)
REGISTRY = []        # (workbook, sheet, title, n_rows, sources)


def check(label, got, expected, tol=0.0):
    ok = (abs(float(got) - float(expected)) <= tol) if not isinstance(expected, str) else (str(got) == expected)
    CHECKS.append((label, got, expected, tol, ok))
    return ok


def clean(v):
    if v is None:
        return None
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return None if (math.isnan(v) or math.isinf(v)) else float(v)
    if isinstance(v, pd.Timestamp):
        return str(v)
    return v


class Book:
    """One workbook; add() writes a panel sheet with the 4-row preamble described in the module docstring."""

    def __init__(self, fname):
        self.fname = fname
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.annot = []

    def add(self, sheet, title, sources, df, desc):
        assert len(sheet) <= 31, sheet
        missing = [c for c in df.columns if c not in desc]
        extra = [c for c in desc if c not in df.columns]
        assert not missing and not extra, (sheet, missing, extra)
        assert len(df) > 0, f"{sheet} is empty"
        ws = self.wb.create_sheet(sheet)
        ws.append([title])
        ws.append(["Source file(s): " + "; ".join(sources)])
        ws.append(list(df.columns))
        ws.append([desc[c] for c in df.columns])
        for row in df.itertuples(index=False, name=None):
            ws.append([clean(v) for v in row])
        ws["A1"].font = Font(bold=True)
        ws["A2"].font = Font(italic=True, color="666666")
        for c in ws[3]:
            c.font = Font(bold=True)
        for c in ws[4]:
            c.font = Font(italic=True, color="666666")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, col in enumerate(df.columns, start=1):
            letter = ws.cell(row=3, column=i).column_letter
            ws.column_dimensions[letter].width = max(14, min(34, len(str(col)) + 4))
        ws.row_dimensions[4].height = 62
        ws.freeze_panes = "A5"
        REGISTRY.append((self.fname, sheet, title, len(df), sources))

    def note(self, panel, quantity, value, note=""):
        self.annot.append({"panel": panel, "quantity": quantity, "value": value, "note": note})

    def save(self, annot_sheet=None, annot_sources=()):
        if annot_sheet and self.annot:
            df = pd.DataFrame(self.annot)
            self.add(annot_sheet, "Values printed on, or used to draw, the figure (annotations, reference lines, axis limits)",
                     list(annot_sources), df,
                     {"panel": "panel letter", "quantity": "what the number is",
                      "value": "value, at full precision unless it is a label", "note": "unit, rule or manuscript cross-reference"})
        path = os.path.join(OUT, self.fname)
        self.wb.save(path)
        return path


def rd(path, **kw):
    return pd.read_csv(path, **kw)


LOW = {c: c.lower() for c in COH}

# =====================================================================================================
# shared inputs for Fig. 1 (make_master_composite.py)
# =====================================================================================================
incr = {c: rd(f"{DD}/merged_incr_{LOW[c]}.csv") for c in COH}
sig = {c: d[(d.fdr_uni < 0.05) & (d.incremental_r2_uni > 0)] for c, d in incr.items()}
counts = {c: len(sig[c]) for c in COH}
tested = {c: int(incr[c]["incremental_r2_uni"].notna().sum()) for c in COH}
medincr = {c: float(sig[c]["incremental_r2_uni"].median()) for c in COH}
seed = {c: rd(f"{DD}/seed_reliability_{LOW[c]}.csv") for c in COH}
for c in COH:
    assert set(seed[c].loc[seed[c]["pinned_significant"], "gene"]) == set(sig[c]["gene"]), c
medheld = {c: float(np.median(seed[c].loc[seed[c]["pinned_significant"],
                                          [f"incr_s{i}" for i in range(1, 10)]].values)) for c in COH}
order_by_count = sorted(COH, key=lambda c: counts[c], reverse=True)
resid = {c: rd(f"{PIN}/{LOW[c]}/residual_results_tumoronly.csv") for c in COH}

# =====================================================================================================
# FIGURE 1
# =====================================================================================================
b1 = Book("Source_Data_Fig1.xlsx")
SRC_INC = [f"figures/figdata/merged_incr_{LOW[c]}.csv" for c in COH]

# ---- 1a: workflow schematic (image panel) ----
t = pd.DataFrame([
    {"panel": "a", "item": "study-design schematic", "file": "figures/Fig1_flagship.png",
     "vector_versions": "figures/Fig1_flagship.svg; figures/Fig1_flagship.pdf; editable source figures/Fig1_concept.drawio",
     "content": "Workflow: protein = mRNA-explained part + residual; H&E slide -> frozen UNI embedding -> incremental R2 across five CPTAC cancers. Drawn artwork; contains no plotted data."},
])
b1.add("Fig1a", "Fig. 1a - study-design schematic (image panel, no numeric data)", ["figures/make_master_composite.py (imports figures/Fig1_flagship.png)"], t,
       {"panel": "panel letter", "item": "what the panel shows", "file": "image file shown", "vector_versions": "vector / editable versions",
        "content": "description of the artwork"})

# ---- 1b: per gene-cohort pair mRNA->protein R2 ----
frames = []
for c in COH:
    d = resid[c]
    frames.append(pd.DataFrame({"cohort": c, "gene": d["gene"], "r2_rna": d["r2_rna"]}))
pair = pd.concat(frames, ignore_index=True)
pair = pair[np.isfinite(pair["r2_rna"])].reset_index(drop=True)
pair["r2_rna_plotted"] = np.clip(pair["r2_rna"], -0.1, 1.0)
pooled = pair["r2_rna"].values
med_pool = float(np.median(pooled))
low_pool = 100 * float(np.mean(pooled < 0.25))
b1.add("Fig1b", "Fig. 1b - out-of-fold mRNA-to-protein R2 for every tested gene-cohort pair (histogram input)",
       [f"server_export/pinned/{LOW[c]}/residual_results_tumoronly.csv" for c in COH], pair,
       {"cohort": "CPTAC cohort", "gene": "gene symbol",
        "r2_rna": "cross-validated R2 of protein on the gene's own mRNA [unitless]",
        "r2_rna_plotted": "same value clipped to [-0.1, 1.0], the range the histogram shows [unitless]"})
hcnt, hedges = np.histogram(np.clip(pooled, -0.1, 1.0), bins=52, range=(-0.1, 1.0))
hb = pd.DataFrame({"bin_left": hedges[:-1], "bin_right": hedges[1:], "n_pairs": hcnt})
b1.add("Fig1b_bins", "Fig. 1b - histogram counts as drawn (52 bins over -0.1 to 1.0, values outside clipped into the end bins)",
       ["computed from sheet Fig1b"], hb,
       {"bin_left": "lower bin edge of mRNA-to-protein R2 [unitless]", "bin_right": "upper bin edge [unitless]",
        "n_pairs": "number of gene-cohort pairs in the bin [count]"})
b1.note("b", "gene-cohort pairs", int(len(pooled)), "finite r2_rna values pooled over the five cohorts; manuscript: 51,553")
b1.note("b", "pooled median R2", med_pool, "dashed line; manuscript: 0.12")
b1.note("b", "% of pairs with R2 < 0.25", low_pool, "shaded region label; manuscript: 67-68%")
b1.note("b", "% of pairs with R2 <= 0", 100 * float(np.mean(pooled <= 0)), "manuscript: 21-26% (per-cohort range)")
b1.note("b", "distinct genes", int(pair["gene"].nunique()), "manuscript: 12,898")
check("Fig1b n gene-cohort pairs", len(pooled), 51553)
check("Fig1b distinct genes", pair["gene"].nunique(), 12898)
check("Fig1b pooled median R2 (2 dp)", round(med_pool, 2), 0.12, 1e-9)
check("Fig1b % pairs R2<0.25 (rounded)", round(low_pool), 67, 1.0)

# ---- 1c: bubble scatter ----
t = pd.DataFrame({
    "cohort": COH, "organ": [ORGAN[c] for c in COH],
    "n_tested_proteins": [tested[c] for c in COH],
    "n_significant_proteins": [counts[c] for c in COH],
    "pct_tested_significant": [100 * counts[c] / tested[c] for c in COH],
    "median_incr_r2_selected": [medincr[c] for c in COH],
    "median_incr_r2_heldout": [medheld[c] for c in COH],
    "bubble_area_pt2": [counts[c] * 0.28 for c in COH]})
b1.add("Fig1c", "Fig. 1c - breadth and effect size per organ (bubble scatter)", SRC_INC + [f"figures/figdata/seed_reliability_{LOW[c]}.csv" for c in COH], t,
       {"cohort": "CPTAC cohort", "organ": "organ",
        "n_tested_proteins": "proteins with a finite UNI incremental R2 [count]",
        "n_significant_proteins": "FDR < 0.05 and incremental R2 > 0 (bubble label) [count]",
        "pct_tested_significant": "y axis: significant / tested [%]",
        "median_incr_r2_selected": "x axis, filled bubble: median incremental R2 of the significant set on the selection seed [unitless]",
        "median_incr_r2_heldout": "x axis, open diamond: median over seeds 1-9 of the same genes (winner's-curse corrected) [unitless]",
        "bubble_area_pt2": "marker area = 0.28 pt2 per significant protein [pt2]"})
for c, want in zip(COH, [2191, 2310, 2566, 702, 1572]):
    check(f"Fig1c significant proteins {c}", counts[c], want)
for c, want in zip(COH, [0.061, 0.032, 0.113, 0.026, 0.040]):
    check(f"Fig1c held-out median {c} (3 dp)", round(medheld[c], 3), want, 1e-9)
b1.note("c", "bubble-size key", "500 and 2,500 proteins", "open grey circles, area 0.28 pt2 per protein")

# ---- 1d: violin data ----
rows = []
for i, c in enumerate(order_by_count, start=1):
    s = sig[c]
    rows.append(pd.DataFrame({"violin_position": i, "cohort": c, "organ": ORGAN[c], "gene": s["gene"].values,
                              "incr_r2_uni": s["incremental_r2_uni"].values, "fdr_uni": s["fdr_uni"].values}))
vd_ = pd.concat(rows, ignore_index=True)
b1.add("Fig1d", "Fig. 1d - incremental R2 of each morphology-predictable protein (violin input)", SRC_INC, vd_,
       {"violin_position": "x position (cohorts ordered by number of significant proteins)", "cohort": "CPTAC cohort", "organ": "organ",
        "gene": "protein / gene symbol", "incr_r2_uni": "incremental R2 of UNI morphology over own mRNA, selection seed [unitless]",
        "fdr_uni": "Benjamini-Hochberg FDR of the permutation test [unitless]"})
allsig = np.concatenate([sig[c]["incremental_r2_uni"].values for c in order_by_count])
b1.note("d", "y-axis upper limit", float(np.percentile(allsig, 99)), "99th percentile of all plotted values; values above are not shown")
for c in COH:
    b1.note("d", f"solid bar, median selected, {c}", medincr[c], "see Fig1c")
    b1.note("d", f"dashed bar, median held-out, {c}", medheld[c], "see Fig1c")

# ---- 1e: UNI vs Phikon ----
frames = []
for c in COH:
    d = incr[c]
    frames.append(pd.DataFrame({"cohort": c, "gene": d["gene"], "incr_r2_uni": d["incremental_r2_uni"], "incr_r2_phikon": d["incremental_r2_phi"]}))
uv = pd.concat(frames, ignore_index=True)
uv = uv[np.isfinite(uv.incr_r2_uni) & np.isfinite(uv.incr_r2_phikon)].reset_index(drop=True)
uv["inside_plotted_window"] = (uv.incr_r2_uni.between(-0.05, 0.35)) & (uv.incr_r2_phikon.between(-0.05, 0.35))
b1.add("Fig1e", "Fig. 1e - per-protein incremental R2, UNI vs Phikon (hexbin input, all tested gene-cohort pairs)", SRC_INC, uv,
       {"cohort": "CPTAC cohort", "gene": "gene symbol", "incr_r2_uni": "incremental R2, UNI encoder [unitless]",
        "incr_r2_phikon": "incremental R2, Phikon encoder [unitless]",
        "inside_plotted_window": "True if both values lie in the plotted window -0.05..0.35 (points outside are not drawn)"})
r_up = float(st.pearsonr(uv.incr_r2_uni, uv.incr_r2_phikon)[0])
b1.note("e", "Pearson r (UNI vs Phikon)", r_up, "all pairs; manuscript: 0.79")
b1.note("e", "n pairs", int(len(uv)), "")
check("Fig1e Pearson r UNI vs Phikon (2 dp)", round(r_up, 2), 0.79, 1e-9)

# ---- 1f: t-SNE ----
emb = rd(f"{DD}/emb_pca50.csv")
Y = TSNE(n_components=2, random_state=0, init="pca", perplexity=30, learning_rate="auto").fit_transform(
    emb[[f"pc{i}" for i in range(1, 51)]].values)
t = pd.DataFrame({"row_in_emb_pca50_csv": np.arange(len(emb)), "cohort": emb["cohort"].str.upper(),
                  "tsne_1": Y[:, 0], "tsne_2": Y[:, 1]})
b1.add("Fig1f", "Fig. 1f - t-SNE of pooled UNI slide embeddings (2,007 slides)", ["figures/figdata/emb_pca50.csv"], t,
       {"row_in_emb_pca50_csv": "0-based row of the slide in emb_pca50.csv (that file carries no slide IDs; slide-level IDs are in the Supplementary Data manifest)",
        "cohort": "CPTAC cohort", "tsne_1": "t-SNE dimension 1 [arbitrary units]", "tsne_2": "t-SNE dimension 2 [arbitrary units]"})
b1.note("f", "t-SNE settings", "PCA(50) input, perplexity 30, init=pca, learning_rate=auto, random_state=0",
        "scikit-learn TSNE; coordinates can differ in the last digits across scikit-learn versions; axes have no units")
check("Fig1f n slides", len(emb), 2007)

# ---- 1g: measured vs morphology-only score ----
frames = []
for c in COH:
    s = rd(f"{DD}/scores_{LOW[c]}.csv", index_col=0)
    frames.append(pd.DataFrame({"cohort": c, "case_id": s.index, "translation_meas": s["translation_meas"].values,
                                "translation_morph": s["translation_morph"].values}))
g1 = pd.concat(frames, ignore_index=True)
b1.add("Fig1g", "Fig. 1g - measured translation-residual score vs morphology-only (out-of-fold) score, per patient",
       [f"figures/figdata/scores_{LOW[c]}.csv" for c in COH], g1,
       {"cohort": "CPTAC cohort", "case_id": "CPTAC case identifier (patient)",
        "translation_meas": "translation-residual score measured from proteomics [z-scored, unitless]",
        "translation_morph": "out-of-fold morphology-only prediction of that score [unitless]"})
xs, ys = g1["translation_meas"].values, g1["translation_morph"].values
rg, pg = st.pearsonr(xs, ys)
b1g, b0g = np.polyfit(xs, ys, 1)
b1.note("g", "Pearson r", float(rg), "manuscript: 0.57")
b1.note("g", "P value", float(pg), "manuscript: 5e-49")
b1.note("g", "OLS slope", float(b1g), "line drawn from min to max of x")
b1.note("g", "OLS intercept", float(b0g), "")
b1.note("g", "n patients", int(len(g1)), "")
check("Fig1g Pearson r (2 dp)", round(rg, 2), 0.57, 1e-9)
check("Fig1g P value (1 significant figure)", f"{pg:.0e}", "5e-49")
b1.save("Fig1_annotations", ["computed from the sheets above"])

# =====================================================================================================
# FIGURE 2  (make_fig_controls.py)
# =====================================================================================================
b2 = Book("Source_Data_Fig2.xlsx")
mg = rd(f"{DD}/merged_incr_ccrcc.csv")
v_all = mg["incremental_r2_uni"].replace([np.inf, -np.inf], np.nan).dropna().values
edges = np.round(np.arange(-0.70, 0.80 + 1e-9, 0.05), 4)
cnt, _ = np.histogram(v_all, bins=edges)
n_below = int((v_all < -0.70).sum())
n_above = int((v_all > 0.80).sum())
median_a = float(np.median(v_all))
frac_pos = float((v_all > 0).mean())
n_sig_a = int(((mg["fdr_uni"] < 0.05) & (mg["incremental_r2_uni"] > 0)).sum())
pct_sig_a = 100.0 * n_sig_a / len(v_all)
t = pd.DataFrame({"bin_left": edges[:-1], "bin_right": edges[1:], "n_genes": cnt,
                  "colour_group": np.where(edges[:-1] >= 0, "positive (orange)", "negative (grey)")})
b2.add("Fig2a", "Fig. 2a - CCRCC incremental R2 of morphology over mRNA, histogram as drawn (0.05 bins)", ["figures/figdata/merged_incr_ccrcc.csv"], t,
       {"bin_left": "lower bin edge [unitless]", "bin_right": "upper bin edge [unitless]", "n_genes": "genes in the bin [count]",
        "colour_group": "bar colour: bins starting at or above 0 are orange"})
gdf = pd.DataFrame({"gene": mg["gene"], "incr_r2_uni": mg["incremental_r2_uni"], "fdr_uni": mg["fdr_uni"]})
gdf = gdf[np.isfinite(gdf.incr_r2_uni)].reset_index(drop=True)
gdf["fdr_significant_positive"] = (gdf.fdr_uni < 0.05) & (gdf.incr_r2_uni > 0)
b2.add("Fig2a_genes", "Fig. 2a - per-gene values behind the histogram (all CCRCC genes with a finite estimate)", ["figures/figdata/merged_incr_ccrcc.csv"], gdf,
       {"gene": "gene symbol", "incr_r2_uni": "incremental R2 of UNI morphology over own mRNA [unitless]",
        "fdr_uni": "Benjamini-Hochberg FDR [unitless]", "fdr_significant_positive": "FDR < 0.05 and incremental R2 > 0"})
cap = rd(f"{PIN}/ccrcc/randproj_control_results.csv")[["incr_real", "incr_gauss", "incr_randproj"]].median()
b2.note("a", "median incremental R2, all genes", median_a, "dashed vertical line; manuscript: -0.14")
b2.note("a", "% of genes with incremental R2 > 0", 100 * frac_pos, "the orange bins (x >= 0) hold this share of all genes")
b2.note("a", "FDR-significant genes", n_sig_a, "manuscript: 2,191")
b2.note("a", "% of genes FDR-significant", pct_sig_a, "manuscript: 22.7%")
b2.note("a", "genes below -0.70 (counted in margin, not drawn)", n_below, "")
b2.note("a", "genes above 0.80 (outside the bin range)", n_above, "")
b2.note("a", "reference line: real WSI PCs", float(cap["incr_real"]), "median incremental R2 (capacity control); manuscript +0.09")
b2.note("a", "reference line: Gaussian noise", float(cap["incr_gauss"]), "manuscript -0.29")
b2.note("a", "reference line: random projection", float(cap["incr_randproj"]), "manuscript -0.08")
check("Fig2a n genes", len(v_all), 9635)
check("Fig2a n FDR-significant", n_sig_a, 2191)
check("Fig2a % FDR-significant (1 dp)", round(pct_sig_a, 1), 22.7, 1e-9)
check("Fig2a median (2 dp)", round(median_a, 2), -0.14, 1e-9)
check("Fig2a real-morphology reference", round(float(cap["incr_real"]), 2), 0.09, 1e-9)
check("Fig2a Gaussian reference", round(float(cap["incr_gauss"]), 2), -0.29, 1e-9)
check("Fig2a random-projection reference", round(float(cap["incr_randproj"]), 2), -0.08, 1e-9)

# ---- 2b, 2c, 2d ----
BL = rd(f"{DD}/batch_levels.csv").set_index("cohort").loc[COH]
assert list(BL.baseline) == [2191, 2310, 2566, 702, 1572]
t = pd.DataFrame({"cohort": COH, "pct_slides_on_dominant_scanner": BL.scan_pct.to_numpy(float),
                  "n_scanners": BL.n_scanners.to_numpy(int), "pct_slides_on_other_scanners": 100 - BL.scan_pct.to_numpy(float)})
b2.add("Fig2b", "Fig. 2b - scanner census (share of slides on the dominant scanner)", ["figures/figdata/batch_levels.csv"], t,
       {"cohort": "CPTAC cohort", "pct_slides_on_dominant_scanner": "x axis: slides on the dominant scanner [%]",
        "n_scanners": "number of scanner IDs among the cohort's slides [count]", "pct_slides_on_other_scanners": "100 minus the dominant share [%]"})
stg = ["baseline", "batch-corrected", "strictest (batch-grouped CV)"]
t = pd.DataFrame([{"cohort": c, "stage_x": i, "stage": s, "n_significant_proteins": int(v)}
                  for c in COH for i, (s, v) in enumerate(zip(stg, [BL.loc[c, "baseline"], BL.loc[c, "batchcorr"], BL.loc[c, "strictest"]]))])
b2.add("Fig2c", "Fig. 2c - significant proteins at baseline, after batch correction and under batch-grouped CV (slope chart)", ["figures/figdata/batch_levels.csv"], t,
       {"cohort": "CPTAC cohort", "stage_x": "x position 0, 1, 2", "stage": "correction level",
        "n_significant_proteins": "FDR < 0.05 and incremental R2 > 0 at that level (y axis, symlog) [count]"})
b2.note("c", "dashed reference line", 50, "50-protein threshold")
C2 = rd(f"{DD}/enrichment_c2_levels.csv")
c2s = C2.groupby(["cohort", "level"])["signature_hit"].sum().unstack("level")
sg = c2s.loc[[c.lower() for c in COH], ["baseline", "batch-corrected", "strictest"]].to_numpy(int)
t = pd.DataFrame([{"cohort": c, "correction_level": lv, "n_signature_family_terms_of_top15": int(sg[i, j])}
                  for i, c in enumerate(COH) for j, lv in enumerate(["baseline", "batch-corrected", "strictest"])])
b2.add("Fig2d", "Fig. 2d - enrichment-family survival: signature-family terms among the top 15 enriched terms", ["figures/figdata/enrichment_c2_levels.csv"], t,
       {"cohort": "CPTAC cohort", "correction_level": "baseline / batch-corrected / strictest",
        "n_signature_family_terms_of_top15": "cell value in the heatmap [count, 0-15]"})
t = C2.copy()
t["cohort"] = t["cohort"].str.upper()
t = t[["cohort", "level", "n_genes", "rank", "term", "qvalue", "overlap", "signature_hit", "families"]]
b2.add("Fig2d_terms", "Fig. 2d - the enriched terms behind each heatmap cell (tested-protein background, top 15 per cohort and level)", ["figures/figdata/enrichment_c2_levels.csv"], t,
       {"cohort": "CPTAC cohort", "level": "correction level", "n_genes": "genes in the significant set at that level [count]",
        "rank": "rank of the term by q value", "term": "Reactome term", "qvalue": "BH-adjusted P of the enrichment [unitless]",
        "overlap": "significant genes in term / term genes tested", "signature_hit": "term falls in one of the six signature families",
        "families": "family label(s)"})

# ---- 2e purity ----
pc = rd(f"{DD}/confound_check_results.csv").replace([np.inf, -np.inf], np.nan).dropna(subset=["incremental_r2", "incr_r2_purity_adjusted", "fdr"])
psig = pc[(pc.fdr < 0.05) & (pc.incremental_r2 > 0) & (pc.incremental_r2 < 1.5) & (pc.incr_r2_purity_adjusted.abs() < 1.5)].copy()
t = pd.DataFrame({"gene": psig.gene.values, "incr_r2_mrna_baseline": psig.incremental_r2.values,
                  "incr_r2_after_purity_adjustment": psig.incr_r2_purity_adjusted.values, "fdr": psig.fdr.values})
b2.add("Fig2e", "Fig. 2e - CCRCC per-gene incremental R2 over mRNA alone vs over mRNA + tumour purity (FDR-significant proteins)", ["figures/figdata/confound_check_results.csv"], t,
       {"gene": "protein / gene symbol", "incr_r2_mrna_baseline": "x axis: incremental R2 over the mRNA-only baseline [unitless]",
        "incr_r2_after_purity_adjustment": "y axis: incremental R2 over mRNA + purity [unitless]", "fdr": "FDR of the mRNA-only test [unitless]"})
px, py = psig.incremental_r2.values, psig.incr_r2_purity_adjusted.values
pr_ = float(st.pearsonr(px, py)[0]); ratio_ = float(np.median(py / px)); pf_ = float(np.mean(py > 0))
b2.note("e", "n proteins plotted", int(len(psig)), "manuscript: 2,191")
b2.note("e", "Pearson r", pr_, "manuscript: 0.85")
b2.note("e", "median retained (median of per-gene ratio, as printed on the panel)", 100 * ratio_, "% ; manuscript: 96%")
b2.note("e", "share remaining > 0", 100 * pf_, "% ; manuscript: 87%")
check("Fig2e n proteins", len(psig), 2191)
check("Fig2e Pearson r (2 dp)", round(pr_, 2), 0.85, 1e-9)
check("Fig2e median retained % (rounded)", round(100 * ratio_), 96, 0.5)
check("Fig2e share remaining >0 % (rounded)", round(100 * pf_), 87, 0.5)

# ---- 2f, 2g, 2h transcriptome-wide baseline ----
tx = rd(f"{DD}/transcriptome_baseline_summary.csv").set_index("cohort").loc[COH]
t = pd.DataFrame({"cohort": COH, "n_genes": tx.n_genes.to_numpy(int), "adv_own_mrna": tx.adv_own.to_numpy(float),
                  "adv_own_mrna_plus_20_rna_pcs": tx.adv_pcs20.to_numpy(float), "retains_pct": tx.retains_pct.to_numpy(float)})
b2.add("Fig2f", "Fig. 2f - advantage of morphology over dimension-matched Gaussian noise, own-mRNA baseline vs + 20 RNA PCs", ["figures/figdata/transcriptome_baseline_summary.csv"], t,
       {"cohort": "CPTAC cohort", "n_genes": "significant proteins evaluated [count]",
        "adv_own_mrna": "x axis: median paired advantage (real minus noise increment), baseline = own mRNA [R2 units]",
        "adv_own_mrna_plus_20_rna_pcs": "y axis: same advantage, baseline = own mRNA + 20 RNA PCs [R2 units]",
        "retains_pct": "label: y/x, retained fraction [%] (as stored in the summary file)"})
frames, pos = [], []
for c in COH:
    gg = rd(f"{DD}/transcriptome_baseline_{LOW[c]}.csv")
    ex = (gg["incr_over_pcs20"] - gg["incr_over_pcs20_randctrl"])
    ok = ex.notna()
    lo, hi = np.percentile(ex[ok].values, [1, 99])
    frames.append(pd.DataFrame({"cohort": c, "gene": gg["gene"][ok].values, "excess_delta_r2": ex[ok].values,
                                "in_central_98pct_drawn": ((ex[ok] >= lo) & (ex[ok] <= hi)).values}))
    pos.append(100 * float((ex[ok] > 0).mean()))
gx = pd.concat(frames, ignore_index=True)
b2.add("Fig2g", "Fig. 2g - per-gene excess of real morphology over same-size noise, transcriptome-wide baseline (violin input)",
       [f"figures/figdata/transcriptome_baseline_{LOW[c]}.csv" for c in COH], gx,
       {"cohort": "CPTAC cohort", "gene": "significant protein / gene symbol",
        "excess_delta_r2": "incremental R2 over own mRNA + 20 RNA PCs, real morphology minus noise control [R2 units]",
        "in_central_98pct_drawn": "True if within the central 98% of the cohort's values (violin is drawn on these only)"})
for c, p in zip(COH, pos):
    b2.note("g", f"% genes with excess > 0 (tick label), {c}", p, "over all genes of the cohort")
t = pd.DataFrame({"cohort": COH, "cc_noise_median": tx.cc_noise_med.to_numpy(float), "cc_morph_median": tx.cc_morph_med.to_numpy(float),
                  "n_morph_cc_gt_0.5_of_20": tx.cc_morph_gt05.to_numpy(int), "n_noise_cc_gt_0.5_of_20": tx.cc_noise_gt05.to_numpy(int)})
b2.add("Fig2h", "Fig. 2h - median canonical correlation with the 20 RNA PCs, morphology (filled) vs same-size noise (open)", ["figures/figdata/transcriptome_baseline_summary.csv"], t,
       {"cohort": "CPTAC cohort", "cc_noise_median": "open marker: median canonical correlation, noise block [unitless]",
        "cc_morph_median": "filled marker: median canonical correlation, 20 morphology PCs [unitless]",
        "n_morph_cc_gt_0.5_of_20": "label: canonical correlations above 0.5, morphology [count of 20]",
        "n_noise_cc_gt_0.5_of_20": "label: canonical correlations above 0.5, noise [count of 20]"})
b2.save("Fig2_annotations", ["computed from the sheets above"])

# ---- Fig. 3a term bookkeeping, copied verbatim from figures/make_fig_biology.py (lines 44-152) so that the same rows are selected ----
Q_SIG, CAP, TOP_N, EXT_N, EXT_Q = 0.05, 30.0, 6, 10, 1e-8
CANON = {
    "Cap-dependent Translation Initiation": [
        "L13a-mediated Translational Silencing Of Ceruloplasmin Expression",
        "GTP Hydrolysis And Joining Of 60S Ribosomal Subunit",
        "Formation Of A Pool Of Free 40S Subunits", "Translation",
        "Response Of EIF2AK4 (GCN2) To Amino Acid Deficiency",
        "Translation Initiation Complex Formation",
        "mRNA Activation Upon Binding Of Cap-Binding Complex And eIFs, Subsequent Binding To 43S",
        "Ribosomal Scanning And Start Codon Recognition",
        "Formation Of Ternary Complex, And Subsequently, 43S Complex",
        "Selenocysteine Synthesis", "Selenoamino Acid Metabolism", "Viral mRNA Translation",
        "Nonsense Mediated Decay (NMD) Independent Of Exon Junction Complex (EJC)",
        "Nonsense Mediated Decay (NMD) Enhanced By Exon Junction Complex (EJC)",
        "Nonsense-Mediated Decay (NMD)", "Eukaryotic Translation Termination",
        "Cellular Response To Starvation", "SARS-CoV-2 Modulates Host Translation Machinery",
        "Influenza Viral RNA Transcription And Replication", "Influenza Infection",
        "Regulation Of Expression Of SLITs And ROBOs", "Signaling By ROBO Receptors"],
    "Eukaryotic Translation Elongation": ["Peptide Chain Elongation"],
    "Major Pathway Of rRNA Processing In Nucleolus And Cytosol": [
        "rRNA Processing In Nucleus And Cytosol", "rRNA Processing",
        "rRNA Modification In Nucleus And Cytosol"],
    "mRNA Splicing - Major Pathway": ["mRNA Splicing"],
    "ER To Golgi Anterograde Transport": ["Transport To Golgi And Subsequent Modification",
                                          "COPI-mediated Anterograde Transport",
                                          "COPII-mediated Vesicle Transport"],
    "Membrane Trafficking": ["Vesicle-mediated Transport"],
    "Transport Of Small Molecules": ["SLC-mediated Transmembrane Transport",
                                     "Disorders Of Transmembrane Transporters"],
    "Chaperonin-mediated Protein Folding": ["Protein Folding"],
    "Prefoldin Mediated Transfer Of Substrate To CCT/TriC": [
        "Cooperation Of Prefoldin And TriC/CCT In Actin And Tubulin Folding",
        "Folding Of Actin By CCT/TriC"],
    "Formation Of Tubulin Folding Intermediates By CCT/TriC": ["Post-chaperonin Tubulin Folding Pathway"],
    "SUMOylation": ["SUMO E3 Ligases SUMOylate Target Proteins"],
    "RHO GTPase Effectors": ["Signaling By Rho GTPases",
                             "Signaling By Rho GTPases, Miro GTPases And RHOBTB3", "RHO GTPase Cycle"],
    "Chromatin Modifying Enzymes": ["HDACs Deacetylate Histones", "Chromatin Organization"],
    "Positive Epigenetic Regulation Of rRNA Expression": [
        "ERCC6 (CSB) And EHMT2 (G9a) Positively Regulate rRNA Expression",
        "Epigenetic Regulation Of Gene Expression", "RNA Polymerase I Transcription",
        "RNA Polymerase I Promoter Clearance", "RNA Polymerase I Transcription Initiation"],
    "Glucose Metabolism": ["Gluconeogenesis"],
    "Assembly Of Collagen Fibrils And Other Multimeric Structures": ["Collagen Formation"],
}
ALIAS = {a: k for k, v in CANON.items() for a in v}
EXCLUDE = {
    "Disease", "Infectious Disease", "Metabolism Of RNA", "Metabolism Of Proteins",
    "Post-translational Protein Modification", "Gene Expression (Transcription)",
    "RNA Polymerase II Transcription", "Cellular Responses To Stress", "Cellular Responses To Stimuli",
    "Innate Immune System", "Immune System", "Metabolism", "Signal Transduction",
    "Protein Localization", "Metabolism Of Amino Acids And Derivatives",
    "Axon Guidance", "Nervous System Development", "Developmental Biology",
    "Interactions Of Vpr With Host Cellular Proteins", "Leishmania Infection", "SARS-CoV Infections",
    "SARS-CoV-2-host Interactions", "Potential Therapeutics For SARS", "Late SARS-CoV-2 Infection Events",
    "Viral Messenger RNA Synthesis", "HIV Infection", "Host Interactions Of HIV Factors",
}
FAMILY = {
    "Cap-dependent Translation Initiation": "translation",
    "Eukaryotic Translation Elongation": "translation",
    "SRP-dependent Cotranslational Protein Targeting To Membrane": "translation",
    "Major Pathway Of rRNA Processing In Nucleolus And Cytosol": "translation",
    "Asparagine N-linked Glycosylation": "secretion",
    "ER To Golgi Anterograde Transport": "secretion",
    "Membrane Trafficking": "secretion",
    "mRNA Splicing - Major Pathway": "splicing",
    "Processing Of Capped Intron-Containing Pre-mRNA": "splicing",
    "Chaperonin-mediated Protein Folding": "folding",
    "Prefoldin Mediated Transfer Of Substrate To CCT/TriC": "folding",
    "Cooperation Of PDCL (PhLP1) And TRiC/CCT In G-protein Beta Folding": "folding",
    "Formation Of Tubulin Folding Intermediates By CCT/TriC": "folding",
    "Extracellular Matrix Organization": "ecm",
    "ECM Proteoglycans": "ecm",
    "Assembly Of Collagen Fibrils And Other Multimeric Structures": "ecm",
    "SUMOylation": "ptm",
    "SUMOylation Of Transcription Cofactors": "ptm",
    "SUMOylation Of DNA Damage Response And Repair Proteins": "ptm",
}
# Row labels are abbreviated so that the widest one fits the ~110 pt label gutter at 7 pt (they
# were tuned for 5.8 pt); the caption spells the Reactome terms out in full.
SHORT = {
    "Cap-dependent Translation Initiation": "Cap-dependent transl. initiation",
    "Eukaryotic Translation Elongation": "Translation elongation",
    "SRP-dependent Cotranslational Protein Targeting To Membrane": "SRP cotranslational targeting",
    "Major Pathway Of rRNA Processing In Nucleolus And Cytosol": "rRNA processing (major)",
    "Asparagine N-linked Glycosylation": "N-linked glycosylation",
    "ER To Golgi Anterograde Transport": "ER-to-Golgi transport",
    "Membrane Trafficking": "Membrane trafficking",
    "mRNA Splicing - Major Pathway": "mRNA splicing (major)",
    "Processing Of Capped Intron-Containing Pre-mRNA": "Processing of capped pre-mRNA",
    "Chaperonin-mediated Protein Folding": "Chaperonin-mediated folding",
    "Prefoldin Mediated Transfer Of Substrate To CCT/TriC": "Prefoldin–CCT/TriC transfer",
    "Cooperation Of PDCL (PhLP1) And TRiC/CCT In G-protein Beta Folding": "PDCL–TRiC/CCT β folding",
    "Formation Of Tubulin Folding Intermediates By CCT/TriC": "Tubulin folding (CCT/TriC)",
    "Extracellular Matrix Organization": "ECM organisation",
    "ECM Proteoglycans": "ECM proteoglycans",
    "Assembly Of Collagen Fibrils And Other Multimeric Structures": "Collagen fibril assembly",
    "SUMOylation": "SUMOylation",
    "SUMOylation Of Transcription Cofactors": "SUMOylation of TF cofactors",
    "SUMOylation Of DNA Damage Response And Repair Proteins": "SUMOylation, DNA repair",
    "Transport Of Small Molecules": "Transport of small molecules",
    "Glucose Metabolism": "Glucose metabolism",
    "Class I Peroxisomal Membrane Protein Import": "Peroxisomal protein import",
    "Neutrophil Degranulation": "Neutrophil degranulation",
    "RHO GTPase Effectors": "RHO GTPase effectors",
    "Chromatin Modifying Enzymes": "Chromatin-modifying enzymes",
    "Positive Epigenetic Regulation Of rRNA Expression": "Epigenetic rRNA activation",
}
FAMS = ["translation", "secretion", "splicing", "folding", "ecm", "ptm", "other"]


FAMTAG = {"translation": "translation", "secretion": "secretion / ER-Golgi", "splicing": "mRNA splicing",
          "folding": "folding / chaperonin", "ecm": "ECM / matrisome", "ptm": "PTM (SUMOylation)", "other": "other"}


def short(t):
    s = SHORT.get(t, t)
    return s if len(s) <= 32 else s[:31] + "…"


def dot_size(q):
    v = min(max(-np.log10(q), 0.0), CAP)
    return 8.0 + 92.0 * v / CAP


# =====================================================================================================
# FIGURE 3  (make_fig_biology.py)
# =====================================================================================================
b3 = Book("Source_Data_Fig3.xlsx")
df = rd(f"{DD}/enrichment_tested_background_full.csv")
df["term"] = df["Term"].str.replace(r"\s+R-HSA-\d+$", "", regex=True)
df["q"] = df["Adjusted P-value"].astype(float)
df["canon"] = df["term"].map(lambda t: ALIAS.get(t, t))
qtab = (df[df["term"] == df["canon"]].pivot_table(index="term", columns="cohort", values="q", aggfunc="min").reindex(columns=COH))


def q_of(term, c):
    return float(qtab.loc[term, c]) if term in qtab.index and pd.notna(qtab.loc[term, c]) else np.nan


picked = {}
for c in COH:
    seen = []
    for _, r in df[df["cohort"] == c].sort_values("q").iterrows():
        k = r["canon"]
        if r["term"] in EXCLUDE or k in EXCLUDE or k in seen:
            continue
        seen.append(k)
        if len(seen) >= EXT_N:
            break
    picked[c] = seen[:TOP_N] + [k for k in seen[TOP_N:EXT_N] if q_of(k, c) < EXT_Q]
rows_ = list(dict.fromkeys(sum(picked.values(), [])))
rows_ = sorted(rows_, key=lambda t: np.nanmin([q_of(t, c) for c in COH]))
fam_of = {t: FAMILY.get(t, "other") for t in rows_}
ordered = []
for f in FAMS:
    ordered += [t for t in rows_ if fam_of[t] == f]
recs = []
for ri, t in enumerate(ordered, start=1):
    for j, c in enumerate(COH, start=1):
        q = q_of(t, c)
        sigf = bool(pd.notna(q) and q < Q_SIG)
        recs.append({"row_top_to_bottom": ri, "reactome_term": t, "row_label_on_figure": short(t), "family": FAMTAG[fam_of[t]],
                     "column_left_to_right": j, "cohort": c, "organ": ORGAN[c],
                     "q_bh": q, "significant_q_lt_0.05": sigf,
                     "neg_log10_q_capped_at_30": (min(max(-np.log10(q), 0.0), CAP) if pd.notna(q) else np.nan),
                     "dot_area_pt2": (dot_size(q) if sigf else 7.0),
                     "dot_style": "filled, family colour" if sigf else "hollow grey (q >= 0.05 or term absent from the cohort's table)"})
t3a = pd.DataFrame(recs)
b3.add("Fig3a", "Fig. 3a - Reactome enrichment dot matrix (tested-proteome background); one row per term x cohort",
       ["figures/figdata/enrichment_tested_background_full.csv"], t3a,
       {"row_top_to_bottom": "row order on the figure (grouped by family, then by smallest q across cohorts)",
        "reactome_term": "Reactome term (nested near-duplicates collapsed as in the figure script)",
        "row_label_on_figure": "abbreviated row label printed on the figure", "family": "signature family (colour band)",
        "column_left_to_right": "column order (CCRCC, LUAD, UCEC, GBM, PDAC)", "cohort": "CPTAC cohort", "organ": "organ",
        "q_bh": "Benjamini-Hochberg adjusted P of the hypergeometric test, minimum over duplicate entries [unitless]; empty if the term is absent",
        "significant_q_lt_0.05": "filled dot if True", "neg_log10_q_capped_at_30": "-log10 q capped at 30 (drives dot area) [unitless]",
        "dot_area_pt2": "marker area: 8 + 92 * min(-log10 q, 30) / 30 for filled dots, 7 for hollow ones [pt2]",
        "dot_style": "how the cell is drawn"})
b3.note("a", "rows on the figure", len(ordered), "terms x 5 cohorts = table rows")
b3.note("a", "rows picked per cohort", "; ".join(f"{c}={len(picked[c])}" for c in COH), "top 6 terms, plus terms 7-10 with q < 1e-8, after exclusions")
b3.note("a", "significance threshold / cap", "q < 0.05; -log10 q capped at 30", "")

vd = rd(f"{DD}/vignette_data.csv")
genes = list(dict.fromkeys(vd["gene"]))
assert genes == ["RPL35A", "SSR3"], genes
letters = {"RPL35A": ("b", "c"), "SSR3": ("d", "e")}


def r2(a, b):
    a, b = np.asarray(a), np.asarray(b)
    return 1 - np.sum((a - b) ** 2) / np.sum((a - a.mean()) ** 2)


VIG = {}
for g in genes:
    d = vd[vd["gene"] == g]
    pl, pr = letters[g]
    t = pd.DataFrame({"patient": d["patient"].values, "mrna_log2_tpm": d["mrna"].values, "protein_measured": d["protein"].values})
    b3.add(f"Fig3{pl}", f"Fig. 3{pl} - {g}: measured protein vs its own mRNA (CCRCC, one point per patient)", ["figures/figdata/vignette_data.csv"], t,
           {"patient": "CPTAC-CCRCC case identifier", "mrna_log2_tpm": "x axis: mRNA abundance [log2 TPM]",
            "protein_measured": "y axis: measured protein abundance [relative units of the input proteomics table]"})
    rr = float(st.pearsonr(d["mrna"], d["protein"])[0]); s1, s0 = np.polyfit(d["mrna"], d["protein"], 1)
    b3.note(pl, f"{g} Pearson r (mRNA vs protein)", rr, "printed on panel")
    b3.note(pl, f"{g} OLS slope / intercept", f"{s1:.6g} / {s0:.6g}", "dashed line, drawn min to max of x")
    VIG[f"{g}_r_mrna"] = rr
    t = pd.DataFrame({"patient": d["patient"].values, "measured_residual": d["residual"].values, "morphology_predicted_residual": d["morph_pred"].values})
    b3.add(f"Fig3{pr}", f"Fig. 3{pr} - {g}: morphology-only prediction of the measured residual (protein minus f(mRNA)), out-of-fold",
           ["figures/figdata/vignette_data.csv"], t,
           {"patient": "CPTAC-CCRCC case identifier", "measured_residual": "x axis: protein minus the mRNA-explained part [relative units]",
            "morphology_predicted_residual": "y axis: out-of-fold prediction from H&E morphology [relative units]"})
    rr2 = float(st.pearsonr(d["residual"], d["morph_pred"])[0]); rsq = float(r2(d["residual"], d["morph_pred"]))
    s1, s0 = np.polyfit(d["residual"], d["morph_pred"], 1)
    b3.note(pr, f"{g} Pearson r (residual vs morphology prediction)", rr2, "printed on panel")
    b3.note(pr, f"{g} morphology R2", rsq, "1 - SS_res/SS_tot of prediction against measured residual; printed on panel")
    b3.note(pr, f"{g} OLS slope / intercept", f"{s1:.6g} / {s0:.6g}", "solid line; dotted line is the identity")
    VIG[f"{g}_r_res"] = rr2
    VIG[f"{g}_R2"] = rsq
check("Fig3b RPL35A r mRNA-protein (2 dp)", round(VIG["RPL35A_r_mrna"], 2), 0.25, 1e-9)
check("Fig3d SSR3 r mRNA-protein (2 dp)", round(VIG["SSR3_r_mrna"], 2), 0.13, 1e-9)
check("Fig3c RPL35A residual r (2 dp)", round(VIG["RPL35A_r_res"], 2), 0.58, 1e-9)
check("Fig3c RPL35A morphology R2 (2 dp)", round(VIG["RPL35A_R2"], 2), 0.23, 1e-9)
check("Fig3e SSR3 residual r (2 dp)", round(VIG["SSR3_r_res"], 2), 0.75, 1e-9)
check("Fig3e SSR3 morphology R2 (2 dp)", round(VIG["SSR3_R2"], 2), 0.51, 1e-9)

# ---- 3f: tile montage (image panel) -> file names and case IDs ----
sc_ = rd(f"{DD}/scores_ccrcc.csv", index_col=0)
rank_ = sc_.dropna(subset=["translation_meas"]).sort_values("translation_meas")
lo_cases, hi_cases = list(rank_.index[:4]), list(rank_.index[-4:])
# the tiles were cut on the server from results/clinical_link_ccrcc_scores.csv (fig_data_export2.py); the local
# copy in the server export must give the same ranking, otherwise the case IDs listed here would be wrong
alt = os.path.join(RES, "results__clinical_link_ccrcc_scores.csv")
a_ = rd(alt, index_col=0).dropna(subset=["translation_meas"]).sort_values("translation_meas")
assert list(a_.index[:4]) == lo_cases and list(a_.index[-4:]) == hi_cases, "tile case ranking differs from server export"
shown_hi = [(1, 0), (2, 0), (2, 1), (3, 1)]
shown_lo = [(0, 1), (1, 1), (3, 0), (3, 1)]
recs = []
for tag, cases, shown, grp in [("hi", hi_cases, shown_hi, "top row: highest measured translation residual"),
                               ("lo", lo_cases, shown_lo, "bottom row: lowest measured translation residual")]:
    for i, case in enumerate(cases):
        for j in (0, 1):
            f = f"figures/figdata/tile_{tag}_{i}_{j}.png"
            col = shown.index((i, j)) + 1 if (i, j) in shown else None
            recs.append({"figure_row": grp, "figure_column": col, "tile_file": f, "case_id": case, "patient_index_in_group": i,
                         "tile_index_within_patient": j, "translation_meas_of_patient": float(sc_.loc[case, "translation_meas"]),
                         "shown_in_figure": col is not None, "file_exists": os.path.exists(os.path.join(HERE, f))})
t3f = pd.DataFrame(recs)
b3.add("Fig3f", "Fig. 3f - H&E tiles, highest vs lowest measured translation-residual CCRCC patients (image panel: file names and case IDs, no pixels)",
       ["figures/figdata/scores_ccrcc.csv", "figures/figdata/tile_hi_*_*.png", "figures/figdata/tile_lo_*_*.png",
        "server_export/scripts/fig_data_export2.py (tile generation rule)"], t3f,
       {"figure_row": "which row of the montage the tile belongs to", "figure_column": "column 1-4 in the montage; empty if the candidate tile is not shown",
        "tile_file": "256 x 256 px tile file (0.4942 um/px, level 0)", "case_id": "CPTAC-CCRCC case the tile was cut from",
        "patient_index_in_group": "0-3, ascending translation residual within the group (the highest patient of the top row is index 3, the lowest patient of the bottom row is index 0)",
        "tile_index_within_patient": "0 or 1: order in which the tile passed the tissue filter",
        "translation_meas_of_patient": "measured translation-residual score of the patient [z-scored, unitless]",
        "shown_in_figure": "True if the tile is drawn in Fig. 3f; all 16 candidate tiles are listed", "file_exists": "tile file present in the repository"})
b3.note("f", "tile size / scale bar", "256 px = 126.5 um at 0.4942 um/px; bar 50 um = 101.2 px", "")
b3.note("f", "tile rule", "first tissue-passing 256-px tiles on a grid starting at (1280,1280) px with stride 2048 px; tissue = >35% of pixels with HSV saturation > 15",
        "server_export/scripts/fig_data_export2.py; per-tile pixel coordinates were not stored")
b3.note("f", "patients per row shown", f"top row {len({i for i, _ in shown_hi})}, bottom row {len({i for i, _ in shown_lo})}",
        "distinct patients among the four tiles shown per row")
b3.save("Fig3_annotations", ["computed from the sheets above"])

# =====================================================================================================
# FIGURE 4  (make_composite2.py)
# =====================================================================================================
b4 = Book("Source_Data_Fig4.xlsx")
lead_term = {"CCRCC": "Chaperonin-mediated Protein Folding", "LUAD": "Cap-dependent Translation Initiation",
             "UCEC": "Asparagine N-linked Glycosylation", "GBM": "mRNA Splicing - Major Pathway",
             "PDAC": "Extracellular Matrix Organization"}
lead_label = {"CCRCC": "Chaperonin protein folding", "LUAD": "Cap-dependent translation", "UCEC": "N-linked glycosylation",
              "GBM": "mRNA splicing (major)", "PDAC": "ECM organisation"}
lead_fam = {"CCRCC": "folding", "LUAD": "translation", "UCEC": "secretion", "GBM": "splicing", "PDAC": "ecm"}
efull = rd(f"{DD}/enrichment_tested_background_full.csv")
recs = []
for j, c in enumerate(COH, start=1):
    row = efull[(efull.cohort == c) & efull.Term.str.startswith(lead_term[c])]
    assert len(row) == 1
    q = float(row["Adjusted P-value"].iloc[0]); v = -np.log10(q)
    recs.append({"column_left_to_right": j, "cohort": c, "organ": ORGAN[c], "reactome_term": row["Term"].iloc[0],
                 "label_on_figure": lead_label[c], "family": FAMTAG[lead_fam[c]], "overlap": row["Overlap"].iloc[0],
                 "q_bh": q, "neg_log10_q": v, "dot_area_pt2": 14.0 + 200.0 * min(max(v, 0.0), 30.0) / 30.0})
t = pd.DataFrame(recs)
b4.add("Fig4a", "Fig. 4a - representative enriched Reactome term per organ", ["figures/figdata/enrichment_tested_background_full.csv"], t,
       {"column_left_to_right": "position on the figure", "cohort": "CPTAC cohort", "organ": "organ", "reactome_term": "full Reactome term name in the enrichment table",
        "label_on_figure": "label printed under the dot", "family": "signature family (dot colour)", "overlap": "significant genes in term / term genes tested",
        "q_bh": "BH-adjusted P (tested-proteome background) [unitless]", "neg_log10_q": "value printed under the label [unitless]",
        "dot_area_pt2": "marker area: 14 + 200 * min(-log10 q, 30) / 30 [pt2]"})

KEY = {"stroma": "incr_W_over_stroma", "immune": "incr_W_over_immune", "proliferation": "incr_W_over_proliferation", "all three": "incr_W_over_ALL"}
HARD = {"stroma": [87, 93, 70, 86, 85], "immune": [99, 78, 92, 77, 90], "proliferation": [86, 32, 83, 44, 48], "all three": [68, 6, 36, 3, 34]}
recs = []
for ai, (ax_, col) in enumerate(KEY.items()):
    for j, c in enumerate(COH):
        d = rd(f"{DD}/composition_controls_{LOW[c]}.csv").dropna(subset=["incr_orig", col])
        mo, ma = float(d["incr_orig"].median()), float(d[col].median())
        pct = 100 * ma / mo
        assert abs(pct - HARD[ax_][j]) <= 1.0
        recs.append({"row": ax_, "cohort": c, "n_genes": len(d), "median_incr_r2_original": mo, "median_incr_r2_adjusted": ma,
                     "retained_pct": pct, "retained_pct_printed": int(round(pct))})
t = pd.DataFrame(recs)
b4.add("Fig4b", "Fig. 4b - share of the morphology increment retained after adding composition scores to the mRNA baseline (heatmap cells)",
       [f"figures/figdata/composition_controls_{LOW[c]}.csv" for c in COH], t,
       {"row": "composition score(s) added to the mRNA baseline", "cohort": "CPTAC cohort", "n_genes": "significant proteins with both estimates [count]",
        "median_incr_r2_original": "median incremental R2 with the mRNA-only baseline [unitless]",
        "median_incr_r2_adjusted": "median incremental R2 after adding the score(s) [unitless]",
        "retained_pct": "100 x median adjusted / median original [%] (ratio of medians, not median of ratios)",
        "retained_pct_printed": "integer printed in the heatmap cell [%]"})

frames = []
for c in COH:
    d = rd(f"{DD}/composition_controls_{LOW[c]}.csv")
    frames.append(pd.concat([pd.DataFrame({"cohort": c}, index=d.index), d], axis=1))
t = pd.concat(frames, ignore_index=True)
b4.add("Fig4b_genes", "Fig. 4b - per-gene values behind the heatmap (significant proteins; original increment and increment after adding each composition score)",
       [f"figures/figdata/composition_controls_{LOW[c]}.csv" for c in COH], t,
       {"cohort": "CPTAC cohort", "gene": "protein / gene symbol", "n": "patients with a measurement [count]",
        "incr_orig": "incremental R2 over the mRNA-only baseline [unitless]",
        "stroma_gain": "change in the baseline's own R2 when the stromal score is added [R2 units]", "incr_W_over_stroma": "incremental R2 over mRNA + stromal score [unitless]",
        "immune_gain": "change in baseline R2 when the immune score is added [R2 units]", "incr_W_over_immune": "incremental R2 over mRNA + immune score [unitless]",
        "proliferation_gain": "change in baseline R2 when the proliferation score is added [R2 units]", "incr_W_over_proliferation": "incremental R2 over mRNA + proliferation score [unitless]",
        "ALL_gain": "change in baseline R2 when all three scores are added [R2 units]", "incr_W_over_ALL": "incremental R2 over mRNA + all three scores [unitless]"})

FEAT = ["nuclear density", "nucleus count", "mean nucleus area", "nucleus area SD", "chromatin OD", "eosin fraction"]
FKEY = ["nuclear_density", "nucleus_count", "mean_nuc_area", "nuc_area_std", "chromatin_od", "eosin_fraction"]
recs = []
for c in COH:
    pc_ = rd(f"{RES}/results__clinical_link_{LOW[c]}_pathomics_corr.csv")
    pc_ = pc_[pc_.score == "translation_meas"].set_index("feature")
    for k, fn in zip(FKEY, FEAT):
        recs.append({"feature": fn, "feature_key": k, "cohort": c, "spearman_rho": float(pc_.loc[k, "rho"]), "p_value": float(pc_.loc[k, "p"])})
t = pd.DataFrame(recs)
b4.add("Fig4c", "Fig. 4c - Spearman correlation of nuclear-morphometry features with the measured translation residual", [f"server_export/results/results__clinical_link_{LOW[c]}_pathomics_corr.csv" for c in COH], t,
       {"feature": "H&E nuclear feature as labelled on the figure", "feature_key": "column name in the source file", "cohort": "CPTAC cohort",
        "spearman_rho": "cell value [unitless]", "p_value": "two-sided uncorrected P of the correlation [unitless]"})

kirc = load_rppa(f"{DD}/kirc_rppa_results.csv")   # 7 modification-antibody genes removed, FDR re-run over the rest
ccr = rd(f"{DD}/merged_incr_ccrcc.csv")
mm = ccr.merge(kirc, on="gene").dropna(subset=["incremental_r2_uni", "incremental_r2"])
mm = mm[(mm.fdr_uni < 0.05) & (mm.incremental_r2_uni > 0)]
okk = (mm.fdr < 0.05) & (mm.incremental_r2 > 0)
kb = rd(f"{REC}/rppa_baserate_all.csv").set_index("cohort").loc["kirc"]
t = pd.DataFrame({"gene": mm.gene.values, "incr_r2_cptac_ccrcc_uni": mm.incremental_r2_uni.values, "fdr_cptac_ccrcc": mm.fdr_uni.values,
                  "incr_r2_tcga_kirc_phikon": mm.incremental_r2.values, "fdr_tcga_kirc": mm.fdr.values, "n_tcga_kirc": mm.n.values,
                  "significant_in_kirc": okk.values})
b4.add("Fig4d", "Fig. 4d - CPTAC-CCRCC morphology-predictable proteins also measured by RPPA in TCGA-KIRC: incremental R2 in both cohorts (slope chart)",
       ["figures/figdata/merged_incr_ccrcc.csv", "figures/figdata/kirc_rppa_results.csv", "review/recalc/rppa_baserate_all.csv"], t,
       {"gene": "protein / gene symbol", "incr_r2_cptac_ccrcc_uni": "left end of the line: incremental R2 in CPTAC-CCRCC, UNI [unitless]",
        "fdr_cptac_ccrcc": "FDR in CPTAC-CCRCC [unitless]", "incr_r2_tcga_kirc_phikon": "right end of the line: incremental R2 in TCGA-KIRC (RPPA), Phikon [unitless]",
        "fdr_tcga_kirc": "FDR in TCGA-KIRC, Benjamini-Hochberg over the 353 proteins that remain after removing the seven modification-antibody genes [unitless]", "n_tcga_kirc": "TCGA-KIRC cases [count]",
        "significant_in_kirc": "FDR < 0.05 and incremental R2 > 0 in KIRC (blue line)"})
b4.note("d", "proteins shown", int(len(mm)), "manuscript: 62")
b4.note("d", "replicated (FDR<0.05, incr > 0 in KIRC)", int(okk.sum()), "manuscript: 35")
b4.note("d", "replicated share", 100 * float(okk.mean()), "% ; manuscript: 56%")
b4.note("d", "median incremental R2, CPTAC-CCRCC (black marker, left)", float(mm.incremental_r2_uni.median()), "")
b4.note("d", "median incremental R2, TCGA-KIRC (black marker, right)", float(mm.incremental_r2.median()), "")
b4.note("d", "panel base rate among other tested RPPA proteins", 100 * float(kb.narrow_rate_base), "% ; manuscript: 47%")
b4.note("d", "Fisher exact P (two-sided)", float(kb.narrow_p_two), "manuscript: 0.20")
b4.note("d", "Fisher odds ratio", float(kb.narrow_OR), "manuscript: 1.46")
check("Fig4d proteins shown", len(mm), 62)
check("Fig4d replicated", int(okk.sum()), 35)
check("Fig4d base rate % (rounded)", round(100 * float(kb.narrow_rate_base)), 47, 0.5)
check("Fig4d Fisher P (2 dp)", round(float(kb.narrow_p_two), 2), 0.20, 1e-9)
check("Fig4d Fisher OR (2 dp)", round(float(kb.narrow_OR), 2), 1.46, 1e-9)

ab = rd(f"{DD}/abmil_ccrcc.csv")
mp_ = ab[ab.method.str.startswith("mean")].iloc[0]
am_ = ab[ab.method.str.startswith("ABMIL")].iloc[0]
t = pd.DataFrame([{"aggregation": mp_.method, "marker": "open", "pearson_r_oof": float(mp_.r), "r2_oof": float(mp_.R2), "n_patients": int(mp_.n)},
                  {"aggregation": am_.method, "marker": "filled", "pearson_r_oof": float(am_.r), "r2_oof": float(am_.R2), "n_patients": int(am_.n)}])
b4.add("Fig4e", "Fig. 4e - aggregation ablation, CCRCC translation residual (same patient folds)", ["figures/figdata/abmil_ccrcc.csv"], t,
       {"aggregation": "tile-embedding aggregation method", "marker": "marker drawn", "pearson_r_oof": "x axis: out-of-fold Pearson r [unitless]",
        "r2_oof": "out-of-fold R2 printed on the panel [unitless]", "n_patients": "patients [count]"})
check("Fig4e ABMIL r (2 dp)", round(float(am_.r), 2), 0.59, 1e-9)
check("Fig4e ABMIL R2 (2 dp)", round(float(am_.R2), 2), 0.33, 1e-9)
check("Fig4e mean-pool r (2 dp)", round(float(mp_.r), 2), 0.41, 1e-9)
check("Fig4e mean-pool R2 (2 dp)", round(float(mp_.R2), 2), 0.14, 1e-9)
b4.save("Fig4_annotations", ["computed from the sheets above"])

# =====================================================================================================
# FIGURE 5  (make_fig_clinical.py)
# =====================================================================================================
b5 = Book("Source_Data_Fig5.xlsx")


def graded_n(cohort):
    """Analysed patients with a gradeable tumour (G1-G4; GX / Unknown / missing excluded).
    NOT clinical_family_all.csv's n column, which is the survival-evaluable n."""
    sc = rd(f"{DD}/scores_{cohort.lower()}.csv", index_col=0)
    cl = rd(f"{DD}/gdc_clinical_{cohort.lower()}.csv").set_index("case").reindex(sc.index)
    return int(cl["tumor_grade"].astype(str).isin(["G1", "G2", "G3", "G4"]).sum())


cfa_all = rd(f"{DD}/clinical_family_all.csv")
cfa = cfa_all[(cfa_all.outcome == "grade") & (cfa_all["mode"] == "H&E-only")]
gc = ["CCRCC", "LUAD", "UCEC", "PDAC"]
recs = []
for fam, lab_ in [("Transl.", "translation"), ("ER-secr.", "ER-secretion")]:
    r = cfa[cfa.family == fam].set_index("cohort").loc[gc]
    n_g = np.array([graded_n(c) for c in gc])
    z = np.arctanh(r["stat"].values); se = 1.0 / np.sqrt(n_g - 3)
    lo, hi = np.tanh(z - 1.96 * se), np.tanh(z + 1.96 * se)
    for k, c in enumerate(gc):
        recs.append({"cohort": c, "score": lab_, "n_patients": int(n_g[k]), "spearman_rho": float(r["stat"].iloc[k]), "ci95_low": float(lo[k]),
                     "ci95_high": float(hi[k]), "p_value": float(r["p"].iloc[k]), "q_bh_cohort_family": float(r["q_cohort"].iloc[k]),
                     "filled_q_lt_0.05": bool(r["q_cohort"].iloc[k] < 0.05)})
t5a = pd.DataFrame(recs)
b5.add("Fig5a", "Fig. 5a - morphology-only residual score vs tumour grade (Spearman rho with 95% Fisher-z interval)", ["figures/figdata/clinical_family_all.csv"], t5a,
       {"cohort": "CPTAC cohort (four gradeable cohorts)", "score": "translation (green) or ER-secretion (orange) residual score, predicted from H&E only",
        "n_patients": "patients with a gradeable tumour (G1-G4) among the analysed patients [count]", "spearman_rho": "dot position [unitless]",
        "ci95_low": "interval lower end, tanh(atanh(rho) - 1.96/sqrt(n-3)) [unitless]", "ci95_high": "interval upper end [unitless]",
        "p_value": "uncorrected two-sided P [unitless]", "q_bh_cohort_family": "BH q within the cohort's own clinical test family [unitless]",
        "filled_q_lt_0.05": "filled dot if True, open otherwise"})

sc = rd(f"{DD}/scores_ccrcc.csv", index_col=0)
cl = rd(f"{DD}/clinical_link_ccrcc_clinical.csv", index_col=0)
dfk = sc.join(cl[["time", "event"]]).dropna(subset=["translation_morph", "time", "event"])
dfk = dfk[dfk["time"] >= 0]
tt = dfk["time"].values.astype(float); ev = dfk["event"].values.astype(int)
gg_ = (dfk["translation_morph"].values >= np.median(dfk["translation_morph"].values)).astype(int)


def km(t, e):
    xs, ys, Sv = [0.0], [1.0], 1.0
    for x in np.unique(t):
        n = (t >= x).sum(); d = int(((t == x) & (e == 1)).sum())
        Sv *= (1 - d / n) if n > 0 else 1
        xs.append(x); ys.append(Sv)
    return np.array(xs), np.array(ys)


def logrank(t, e, g):
    O1 = E1 = V = 0.0
    for x in np.unique(t[e == 1]):
        risk = t >= x; n = risk.sum(); n1 = (risk & (g == 1)).sum()
        d = (t == x) & (e == 1); dd = int(d.sum()); d1 = int((d & (g == 1)).sum())
        O1 += d1
        if n > 1:
            E1 += dd * n1 / n; V += dd * (n1 / n) * (1 - n1 / n) * (n - dd) / (n - 1)
    return float(1 - chi2.cdf((O1 - E1) ** 2 / V, 1)) if V > 0 else 1.0


pv = logrank(tt, ev, gg_)
t = pd.DataFrame({"case_id": dfk.index, "translation_morph_score": dfk["translation_morph"].values, "group": np.where(gg_ == 1, "high score", "low score"),
                  "time_days": tt, "event_death": ev})
b5.add("Fig5b", "Fig. 5b - CCRCC overall survival by median split of the morphology-only translation-residual score (patient-level input)",
       ["figures/figdata/scores_ccrcc.csv", "figures/figdata/clinical_link_ccrcc_clinical.csv"], t,
       {"case_id": "CPTAC-CCRCC case identifier", "translation_morph_score": "out-of-fold morphology-only translation-residual score [unitless]",
        "group": "high = at or above the cohort median of the score", "time_days": "overall-survival time, death or last follow-up [days]",
        "event_death": "1 = died, 0 = censored"})
recs = []
for gi, name in [(1, "high score"), (0, "low score")]:
    m = gg_ == gi
    x, y = km(tt[m], ev[m])
    for a, b in zip(x, y):
        recs.append({"group": name, "n_in_group": int(m.sum()), "time_days": float(a), "survival_probability": float(b)})
t = pd.DataFrame(recs)
b5.add("Fig5b_curves", "Fig. 5b - Kaplan-Meier step coordinates as drawn (steps 'post')", ["computed from sheet Fig5b"], t,
       {"group": "score group", "n_in_group": "patients in the group [count]", "time_days": "step time [days]", "survival_probability": "survival probability after the step [unitless]"})
b5.note("b", "log-rank P", pv, "descriptive; manuscript: 0.031")
b5.note("b", "patients analysed", int(len(dfk)), "manuscript: 98")
b5.note("b", "deaths", int(ev.sum()), "")
b5.note("b", "n high / n low", f"{int((gg_ == 1).sum())} / {int((gg_ == 0).sum())}", "legend on panel")
check("Fig5b log-rank P (3 dp)", round(pv, 3), 0.031, 1e-9)
check("Fig5b patients", len(dfk), 98)

cox = json.load(open(f"{REC}/recalc_survival_raw.json", encoding="utf-8"))["cox"]


def _term(block, name="translation_morph"):
    x = next(x for x in cox[block]["terms"] if x["term"] == name)
    return x["HR"], x["p"], x["CI95"]


recs = []
for nm, blk in [("univariable", "uni_score"), ("adj. grade+stage", "adj_grade_stage")]:
    hr, p, ci = _term(blk)
    recs.append({"model": nm, "hazard_ratio_per_sd": float(hr), "ci95_low": float(ci[0]), "ci95_high": float(ci[1]), "p_value": float(p),
                 "n_patients": int(cox[blk]["n"]), "n_events": int(cox[blk]["events"]) if cox[blk].get("events") is not None else np.nan})
t5c = pd.DataFrame(recs)
b5.add("Fig5c", "Fig. 5c - Cox hazard ratio of the morphology-only translation score, univariable vs adjusted for grade and stage", ["review/recalc/recalc_survival_raw.json"], t5c,
       {"model": "Cox model", "hazard_ratio_per_sd": "point (per SD of the score) [unitless]", "ci95_low": "model 95% CI lower [unitless]", "ci95_high": "model 95% CI upper [unitless]",
        "p_value": "Wald P of the score term [unitless]", "n_patients": "patients in the model [count]", "n_events": "deaths [count]"})
check("Fig5c univariable HR (2 dp)", round(float(t5c.hazard_ratio_per_sd[0]), 2), 1.71, 1e-9)
check("Fig5c univariable p (3 dp)", round(float(t5c.p_value[0]), 3), 0.007, 1e-9)
check("Fig5c adjusted HR (2 dp)", round(float(t5c.hazard_ratio_per_sd[1]), 2), 1.33, 1e-9)
check("Fig5c adjusted p vs the '0.269' printed in the Fig. 5 caption (3 dp)", round(float(t5c.p_value[1]), 3), 0.269, 1e-9)
b5.note("c", "adjusted p, unrounded", float(t5c.p_value[1]), "the figure prints 0.269 (3 dp)")
b5.save("Fig5_annotations", ["computed from the sheets above"])

# =====================================================================================================
# MAIN TABLES 1 and 2
# =====================================================================================================
bt = Book("Source_Data_Tables.xlsx")
fdr_only = {c: int((resid[c].fdr < 0.05).sum()) for c in COH}
patients = {c: len(rd(f"{DD}/scores_{LOW[c]}.csv")) for c in COH}
t = pd.DataFrame({"cohort": COH, "organ": [ORGAN[c] for c in COH], "patients": [patients[c] for c in COH],
                  "sig_proteins": [counts[c] for c in COH], "sig_pct_of_tested": [100 * counts[c] / tested[c] for c in COH],
                  "n_tested_proteins": [tested[c] for c in COH],
                  "heldout_median_incr_r2": [medheld[c] for c in COH], "selection_seed_median_incr_r2": [medincr[c] for c in COH],
                  "batch_corrected_sig": BL.batchcorr.to_numpy(int), "strictest_sig": BL.strictest.to_numpy(int),
                  "dominant_scanner_pct": BL.scan_pct.to_numpy(float), "fdr_only_without_positivity_filter": [fdr_only[c] for c in COH]})
bt.add("Table1", "Table 1 - five-cohort atlas",
       SRC_INC + ["figures/figdata/seed_reliability_<cohort>.csv", "figures/figdata/batch_levels.csv", "figures/figdata/scores_<cohort>.csv",
                  "server_export/pinned/<cohort>/residual_results_tumoronly.csv (FDR-only counts)"], t,
       {"cohort": "CPTAC cohort", "organ": "organ", "patients": "patients with a score file entry [count]",
        "sig_proteins": "Sig.: FDR < 0.05 and incremental R2 > 0 [count]", "sig_pct_of_tested": "Sig. (%) [% of tested proteins]",
        "n_tested_proteins": "tested proteins [count]",
        "heldout_median_incr_r2": "Held-out R2: median over seeds 1-9 of the Sig. set [unitless]",
        "selection_seed_median_incr_r2": "median on the selection seed (caption: 0.054-0.118) [unitless]",
        "batch_corrected_sig": "Batch-corr. [count]", "strictest_sig": "Strictest (batch-grouped CV) [count]",
        "dominant_scanner_pct": "Dom. scanner % [%]", "fdr_only_without_positivity_filter": "caption: FDR < 0.05 without the positivity filter [count]"})
for c, want in zip(COH, [2904, 3086, 3263, 714, 1844]):
    check(f"Table1 FDR-only count {c}", fdr_only[c], want)
for c, want in zip(COH, [103, 105, 100, 99, 137]):
    check(f"Table1 patients {c}", patients[c], want)
for c, want in zip(COH, [815, 726, 201, 680, 682]):
    check(f"Table1 batch-corrected {c}", int(BL.loc[c, "batchcorr"]), want)
for c, want in zip(COH, [1178, 802, 253, 579, 342]):
    check(f"Table1 strictest {c}", int(BL.loc[c, "strictest"]), want)
for c, want in zip(COH, [99.2, 98.6, 99.0, 100.0, 92.1]):
    check(f"Table1 dominant scanner % {c}", float(BL.loc[c, "scan_pct"]), want, 1e-9)
for c, want in zip(COH, [22.7, 21.5, 24.4, 6.5, 15.9]):
    check(f"Table1 sig % {c} (1 dp)", round(100 * counts[c] / tested[c], 1), want, 1e-9)
check("Table1 selection-seed median range low (3 dp)", round(min(medincr.values()), 3), 0.054, 1e-9)
check("Table1 selection-seed median range high (3 dp)", round(max(medincr.values()), 3), 0.118, 1e-9)

cf = cfa_all[cfa_all["mode"] == "H&E-only"]


def cfv(cohort, fam, outcome, col):
    r = cf[(cf.cohort == cohort) & (cf.family == fam) & (cf.outcome == outcome)]
    return float(r[col].iloc[0]) if len(r) else np.nan


recs = []
for c in ["CCRCC", "LUAD", "PDAC", "UCEC", "GBM"]:
    recs.append({"cohort": c,
                 "n_graded": (graded_n(c) if c != "GBM" else np.nan), "n_survival_evaluable_csv": cfv(c, "Transl.", "grade", "n"), "n_survival_test": cfv(c, "Transl.", "survival (log-rank)", "n"),
                 "transl_rho_grade": cfv(c, "Transl.", "grade", "stat"), "transl_rho_grade_q_cohort": cfv(c, "Transl.", "grade", "q_cohort"),
                 "ersecr_rho_grade": cfv(c, "ER-secr.", "grade", "stat"), "ersecr_rho_grade_q_cohort": cfv(c, "ER-secr.", "grade", "q_cohort"),
                 "transl_rho_stage": cfv(c, "Transl.", "stage", "stat"), "transl_rho_stage_q_cohort": cfv(c, "Transl.", "stage", "q_cohort"),
                 "transl_survival_p": cfv(c, "Transl.", "survival (log-rank)", "p"), "transl_survival_q_cohort": cfv(c, "Transl.", "survival (log-rank)", "q_cohort")})
t2 = pd.DataFrame(recs)
bt.add("Table2", "Table 2 - morphology-only residual score vs grade, stage and survival (H&E-only mode)", ["figures/figdata/clinical_family_all.csv"], t2,
       {"cohort": "CPTAC cohort (row order as in the table)",
        "n_graded": "analysed patients with a gradeable tumour (G1-G4) [count]; empty for GBM (no grade or stage recorded)",
        "n_survival_evaluable_csv": "the 'n' column of clinical_family_all.csv: the survival-evaluable n, NOT the grade n [count]",
        "n_survival_test": "n of the survival test (Table 2 in parentheses) [count]",
        "transl_rho_grade": "Transl. rho_grade [Spearman, unitless]", "transl_rho_grade_q_cohort": "BH q within the cohort family (drives the stars) [unitless]",
        "ersecr_rho_grade": "ER-secr. rho_grade [Spearman, unitless]", "ersecr_rho_grade_q_cohort": "BH q within the cohort family [unitless]",
        "transl_rho_stage": "Transl. rho_stage [Spearman, unitless]", "transl_rho_stage_q_cohort": "BH q within the cohort family [unitless]",
        "transl_survival_p": "Transl. surv. p, median-split log-rank, uncorrected [unitless]", "transl_survival_q_cohort": "BH q of the survival test [unitless]"})
TAB2 = {"CCRCC": (0.37, 0.44, 0.33, 0.031), "LUAD": (0.40, 0.25, 0.15, 0.17), "PDAC": (0.21, 0.19, -0.03, 0.20),
        "UCEC": (0.16, -0.07, -0.03, 0.72), "GBM": (np.nan, np.nan, np.nan, 0.47)}
for _, r in t2.iterrows():
    w = TAB2[r.cohort]
    for lab_, got, want in [("Transl. rho_grade", r.transl_rho_grade, w[0]), ("ER-secr. rho_grade", r.ersecr_rho_grade, w[1]),
                            ("Transl. rho_stage", r.transl_rho_stage, w[2])]:
        if not np.isnan(want):
            check(f"Table2 {r.cohort} {lab_} (2 dp)", round(got, 2), want, 1e-9)
    check(f"Table2 {r.cohort} Transl. surv. p (2-3 dp as printed)", round(r.transl_survival_p, 3 if r.cohort == "CCRCC" else 2), w[3], 1e-9)
TAB2_N = {"CCRCC": (103, 98), "LUAD": (103, None), "PDAC": (137, None), "UCEC": (97, None), "GBM": (99, None)}
for _, r in t2.iterrows():
    n_print = TAB2_N[r.cohort][0]
    got = r.n_survival_test if r.cohort == "GBM" else r.n_graded
    check(f"Table2 {r.cohort} n printed vs clinical_family_all.csv", got, n_print)
bt.save()

# =====================================================================================================
# LONG SUPPLEMENTARY TABLES
# =====================================================================================================
bs = Book("Supplementary_Tables_long.xlsx")
core = rd(f"{DD}/pan_organ_core.csv", index_col=0)
core.index.name = "gene"
mach_nice = {"translocon_OST_SPC": "Translocon/OST/SPC", "spliceosome_mRNA_export": "Spliceosome/export", "cohesin": "Cohesin",
             "ribosome_translation": "Ribosome/translation", "other": "Unclassified"}
tex = open(os.path.join(HERE, "SuppTable_core.tex"), encoding="utf-8").read().replace("\r\n", "\n")
gene_order = []
for line in tex.split("\n"):
    parts = [p.strip() for p in line.rstrip("\\ ").split("&")]
    if len(parts) == 8 and parts[1] in core.index:
        gene_order.append(parts[1])
assert len(gene_order) == len(core) == 170 and len(set(gene_order)) == 170, (len(gene_order), len(core))
core = core.loc[gene_order]
t = pd.DataFrame({"row_in_supp_table": np.arange(1, len(core) + 1), "machine": [mach_nice[m] for m in core["machine"]], "gene": core.index,
                  **{f"incr_r2_{c}": core[c.lower()].values for c in COH}, "k_cohorts_significant": core["k_of_5"].astype(int).values})
bs.add("SuppTable_core", "Supplementary Table (core) - 170 genes FDR-significant with positive incremental R2 in >= 4 of 5 cohorts (rows in the order of the printed table)",
       ["figures/figdata/pan_organ_core.csv (rounded to 4 dp as stored)", "SuppTable_core.tex (row order)", "pan_organ_core.py"], t,
       {"row_in_supp_table": "row order in the printed table", "machine": "obligate multiprotein machine where classifiable, otherwise 'Unclassified'", "gene": "gene symbol",
        **{f"incr_r2_{c}": f"incremental R2 of morphology over own mRNA, UNI, {c} [unitless; printed to 2 dp with sign]" for c in COH},
        "k_cohorts_significant": "number of cohorts (of 5) in which the gene is FDR-significant with positive incremental R2 [count]"})
# rounding check against the printed table
bad = 0
for line in tex.split("\n"):
    parts = [p.strip() for p in line.rstrip("\\ ").split("&")]
    if len(parts) == 8 and parts[1] in core.index:
        for c, s in zip(COH, parts[2:7]):
            v = core.loc[parts[1], c.lower()]
            if (s == "--" and pd.notna(v)) or (s != "--" and abs(float(s) - float(v)) > 0.0051):
                bad += 1
check("SuppTable_core cells that disagree with the printed 2-dp value", bad, 0)
check("SuppTable_core genes", len(core), 170)
check("SuppTable_core classified genes", int((core["machine"] != "other").sum()), 50)

d = rd(f"{DD}/clinical_family_all.csv")
ORD = {"CCRCC": 0, "LUAD": 1, "UCEC": 2, "GBM": 3, "PDAC": 4}
d = d.sort_values(["cohort", "family", "mode", "outcome"], key=lambda s: s.map(ORD) if s.name == "cohort" else s, kind="mergesort").reset_index(drop=True)
t = pd.DataFrame({"row_in_supp_table": np.arange(1, len(d) + 1), "cohort": d.cohort, "score_family": d.family, "input": d["mode"], "outcome": d.outcome,
                  "statistic": d.stat, "p_value": d.p, "n_patients": d.n, "adjustment_flag": d.added, "q_within_cohort": d.q_cohort, "q_over_all_108": d.q_all})
bs.add("SuppTable_clinical_family", "Supplementary Table (clinical family) - all 108 clinical tests", ["figures/figdata/clinical_family_all.csv", "make_supptable_clinical_family.py (row order)"], t,
       {"row_in_supp_table": "row order in the printed table", "cohort": "CPTAC cohort", "score_family": "residual pathway score: Transl. / ER-secr. / Matrisome",
        "input": "meas. = measured proteomic score; H&E-only = morphology-predicted score", "outcome": "grade, stage, survival (log-rank) or Cox model",
        "statistic": "Spearman rho (grade, stage), log-rank chi-square (survival) or hazard ratio (Cox) [unitless]", "p_value": "uncorrected P [unitless]",
        "n_patients": "patients in the test [count]", "adjustment_flag": "'age' = age-adjusted Cox (footnote section sign); 'missing-adjusted' = grade+stage-adjusted Cox for the CCRCC ER-secretion and matrisome scores (footnote double dagger)",
        "q_within_cohort": "BH q within the cohort's own family (primary correction) [unitless]", "q_over_all_108": "BH q over all 108 tests [unitless]"})
check("SuppTable_clinical_family rows", len(d), 108)
check("SuppTable_clinical_family age-adjusted rows", int((d.added == "age").sum()), 22)
for c, want in zip(COH, [30, 20, 20, 8, 30]):
    check(f"SuppTable_clinical_family tests {c}", int((d.cohort == c).sum()), want)

tp = rd(os.path.join(HERE, "SuppTable_top_proteins.csv"))
t = pd.DataFrame({"row_in_supp_table": np.arange(1, len(tp) + 1), "cohort": tp.cohort, "protein": tp.gene, "n_patients": tp.n, "incr_r2": tp.incr_r2, "fdr": tp.fdr, "family": tp.family})
bs.add("SuppTable_top_proteins", "Supplementary Table (top proteins) - the 12 FDR-significant proteins with the largest incremental R2 per cohort", ["SuppTable_top_proteins.csv", "figures/make_top_table.py"], t,
       {"row_in_supp_table": "row order in the printed table", "cohort": "CPTAC cohort", "protein": "protein / gene symbol", "n_patients": "patients with a measurement [count]",
        "incr_r2": "incremental R2 of morphology over own mRNA, UNI [unitless; printed to 2 dp]", "fdr": "Benjamini-Hochberg FDR [unitless; printed to 2 significant figures]",
        "family": "functional annotation"})
bad = 0
for c in COH:
    s = resid[c]
    s = s[(s.fdr < 0.05) & (s.incremental_r2 > 0)].sort_values("incremental_r2", ascending=False).head(12)
    mine = tp[tp.cohort == c].reset_index(drop=True)
    bad += int(list(s.gene) != list(mine.gene)) + int(np.abs(s.incremental_r2.values - mine.incr_r2.values).max() > 1e-6)
check("SuppTable_top_proteins cohorts whose top-12 list disagrees with server_export/pinned", bad, 0)
check("SuppTable_top_proteins rows", len(tp), 60)
bs.save()


# =====================================================================================================
# cross-check report + README
# =====================================================================================================
lines = ["Cross-check of plotted / tabulated numbers against the manuscript text (main.tex), generated by make_source_data.py", ""]
n_bad = 0
for label, got, exp, tol, ok in CHECKS:
    g = f"{got:.6g}" if isinstance(got, (float, np.floating)) else str(got)
    lines.append(f"[{'OK  ' if ok else 'DIFF'}] {label}: computed {g}; manuscript {exp}")
    n_bad += (not ok)
lines += ["", f"{len(CHECKS) - n_bad} of {len(CHECKS)} checks agree; {n_bad} differ."]
with open(os.path.join(OUT, "crosscheck_report.txt"), "w", encoding="utf-8") as f:
    f.write("\n".join(lines) + "\n")

WB_DESC = {
    "Source_Data_Fig1.xlsx": "Fig. 1 (make_master_composite.py): design schematic, mRNA-to-protein R2 distribution, breadth and effect size per organ, incremental-R2 violins, UNI vs Phikon, t-SNE of slide embeddings, measured vs morphology-only score.",
    "Source_Data_Fig2.xlsx": "Fig. 2 (make_fig_controls.py): CCRCC incremental-R2 histogram with capacity controls, scanner census, batch-retention slope chart, enrichment-family survival, tumour-purity control, transcriptome-wide baseline (three panels).",
    "Source_Data_Fig3.xlsx": "Fig. 3 (make_fig_biology.py): Reactome enrichment dot matrix, the RPL35A and SSR3 single-protein vignettes, and the H&E tile montage (case IDs and file names).",
    "Source_Data_Fig4.xlsx": "Fig. 4 (make_composite2.py): representative enriched term per organ, composition controls, nuclear-morphometry correlates, TCGA-KIRC replication slope chart, ABMIL vs mean-pooling.",
    "Source_Data_Fig5.xlsx": "Fig. 5 (make_fig_clinical.py): residual score vs grade, CCRCC Kaplan-Meier, Cox hazard ratios.",
    "Source_Data_Tables.xlsx": "Main Tables 1 and 2 with the extra columns (q values, unrounded values) that the printed tables round or omit.",
    "Supplementary_Tables_long.xlsx": "The three long Supplementary Tables (core genes, complete clinical test family, top proteins), from the CSVs they were generated from.",
}
readme = ["# Source Data", "",
          "Source Data for the MorphoResidual manuscript (Nature Communications submission). Every workbook is regenerated by `../make_source_data.py` (run `python make_source_data.py` "
          "from the project folder); the script recomputes each panel from the same input files the figure scripts read and imports nothing "
          "from them.", "",
          "## Files", "", "| File | Contents |", "|---|---|"]
for k, v in WB_DESC.items():
    readme.append(f"| `{k}` | {v} |")
readme += ["| `crosscheck_report.txt` | Numbers recomputed here and compared with the manuscript text (counts, r, P, HR, medians). |",
           "| `README.md` | This file. |", "",
           "## Sheet layout", "",
           "One workbook per main figure, one sheet per panel (`Fig1a`, `Fig1b`, ...). Panels that need a second table (for example the raw per-gene "
           "values behind a histogram) have an extra sheet with a suffix (`Fig1b_bins`, `Fig2a_genes`, `Fig5b_curves`). A final sheet `FigN_annotations` "
           "lists the numbers printed on the figure or used to draw reference lines and axis limits. Every sheet has the same layout:", "",
           "| Row | Content |", "|---|---|",
           "| 1 | panel title |", "| 2 | source file(s), relative to the project root |",
           "| 3 | column names (units in the explanation row) |", "| 4 | one-line explanation and unit of each column |",
           "| 5 onward | data |", "",
           "Read a sheet in Python with `pandas.read_excel(path, sheet_name='Fig1c', header=2, skiprows=[3])`.", "",
           "## Notes", "",
           "* Image panels (Fig. 1a schematic, Fig. 3f H&E tiles) carry file names and case identifiers instead of pixels. Fig. 3f lists all 16 candidate tiles "
           "and marks the eight shown; per-tile pixel coordinates were not stored, the rule that generated them is given in the sheet `Fig3_annotations`.",
           "* Panels that plot a smoothed or binned summary (histograms, hexbin, violins) have the underlying per-gene or per-patient values plus, where useful, the bin counts.",
           "* The t-SNE coordinates (Fig. 1f) are recomputed with the figure script's settings (`random_state=0`); the last digits can differ across scikit-learn versions. "
           "The embedding file carries no slide identifiers; slide-, patient- and batch-level identifiers are in the Supplementary Data manifest.",
           "* Values are stored at full precision; the figures and tables print rounded values.",
           "* No whole-slide images, proteomic spectra or raw sequence data are included (Data availability).", "",
           "## Sheets", "", "| Workbook | Sheet | Rows | Title |", "|---|---|---|---|"]
for wbn, sh, title, n, srcs in REGISTRY:
    readme.append(f"| `{wbn}` | `{sh}` | {n:,} | {title.replace('|', '/')} |")
readme += ["", "## Cross-check summary", "", f"{len(CHECKS) - n_bad} of {len(CHECKS)} recomputed numbers agree with the manuscript text; "
           f"{n_bad} differ (listed as DIFF in `crosscheck_report.txt`).", ""]
with open(os.path.join(OUT, "README.md"), "w", encoding="utf-8") as f:
    f.write("\n".join(readme))

print(f"wrote {len(set(r[0] for r in REGISTRY))} workbooks, {len(REGISTRY)} sheets to {OUT}")
print(f"cross-checks: {len(CHECKS) - n_bad}/{len(CHECKS)} agree")
for label, got, exp, tol, ok in CHECKS:
    if not ok:
        print("  DIFF:", label, "| computed", got, "| manuscript", exp)
